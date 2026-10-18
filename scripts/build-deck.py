"""Fills the organisers' AI Builder Cup template with the Scam Shield content.

Usage (from the repo root):
  python scripts/build-deck.py <template.pptx> <out.pptx>

The template is docs/submission/template-base.pptx: the organisers' own file with every shape that an earlier
version of this script added removed again (only the shapes named "Google Shape;..." remain, plus the original
slide titles). The script never changes that file, so running it twice always gives the same deck.
Benchmark numbers are read from eval/results/latest.md (and run-hindi.md when it exists) so the deck cannot drift
from the evidence. All 16 template slides are kept in order, with their branded header and footer.
Needs python-pptx, Pillow and qrcode (pip install qrcode in the same environment).
"""
import datetime
import os
import re
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFilter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
SUB = ROOT / "docs" / "submission"
SHOTS = SUB / "screenshots"
IMG = SUB / "deck-images"
GEN = IMG / "gen"  # images this script composes (phone frames, QR); rebuilt on every run
EVAL = ROOT / "eval" / "results"

# Avenir Next ships with macOS and PowerPoint embeds it in the PDF; the template's own Google Sans Flex is not installed here.
FONT = "Avenir Next"
FONT_MED = "Avenir Next Medium"
FONT_BOLD = "Avenir Next Demi Bold"


def rgb(h):
    return RGBColor.from_string(h.lstrip("#").upper())


NAVY = rgb("#081B32")
NAVY2 = rgb("#102B50")
BLUE = rgb("#1764E8")
CYAN = rgb("#34C9E8")
TINT = rgb("#E8F0FE")
TINT2 = rgb("#F1F5FB")
INK = rgb("#0B1F33")
MUTED = rgb("#3F4B5B")  # darker than a typical grey: 8.9:1 on white, 7.9:1 on the light tints
BORDER = rgb("#D5DEEB")
WHITE = rgb("#FFFFFF")
RED = rgb("#B42318")
RED_TINT = rgb("#FDECEB")
AMBER = rgb("#9A5B00")
AMBER_TINT = rgb("#FFF4DB")
SLATE = rgb("#475467")
SLATE_TINT = rgb("#EEF2F6")

REPO_URL = "https://github.com/Namans12/scam-shield"
LIVE_URL = "https://project-91260b43-6705-4176-b71.web.app"
VIDEO_PLACEHOLDER = "[add the public video link here]"
VIDEO = os.environ.get("DECK_VIDEO_URL", VIDEO_PLACEHOLDER)  # DECK_VIDEO_URL=https://... scripts/build-deck.sh
TRACK = "BFSI: Intelligent Risk, Fraud & Financial Experiences"

MX = 0.4  # left and right margin (inches); content width 9.2
CW = 9.2
Y0 = 1.32  # content top (below the template's title)
Y1 = 5.38  # content bottom (above the template's footer bar)


# ---------- benchmark numbers ----------
def read_results(path: Path) -> dict:
    text = path.read_text(encoding="utf-8")
    pick = lambda label: re.search(rf"\|\s*{re.escape(label)}[^|]*\|\s*([^|]+?)\s*\|", text).group(1)
    lat = {m.group(1): (m.group(2), m.group(3), m.group(4)) for m in re.finditer(r"\|\s*(text|image|audio)\s*\|\s*([^|]+?)\s*\|\s*([^|]+?)\s*\|\s*(\d+)\s*\|", text)}
    return {
        "accuracy": pick("Accuracy"),
        "recall": pick("Scam recall"),
        "fpr": pick("False-positive rate"),
        "type": pick("Scam type correct"),
        "errors": pick("Errors"),
        "latency": lat,
        "cases": re.search(r"(\d+) synthetic cases: (\d+) scams, (\d+) genuine", text).groups(),
        "false_positives": re.search(r"False positives: (.+)", text).group(1).strip(),
        "date": re.search(r"Run: (\d{4}-\d{2}-\d{2})", text).group(1),
        "lang": re.search(r"language: (\w+)", text).group(1),
    }


# ---------- low-level helpers ----------
class P:
    """A paragraph: runs = list of (text, opts) or str; opts: size, font ('b' demi, 'm' medium), color, italic."""

    def __init__(self, runs, bullet=None, after=0, before=0, align=None, line=1.0, bcolor=None, hang=0.17):
        self.runs = [(r, {}) if isinstance(r, str) else r for r in (runs if isinstance(runs, list) else [runs])]
        self.bullet, self.after, self.before, self.align, self.line, self.bcolor, self.hang = bullet, after, before, align, line, bcolor, hang


def style_run(run, size, font=None, color=INK, italic=False):
    run.font.name = {"b": FONT_BOLD, "m": FONT_MED}.get(font, FONT)
    run.font.size = Pt(size)
    run.font.italic = italic
    run.font.color.rgb = color


def tx(slide, x, y, w, h, paras, size=12, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, name=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, para in enumerate(paras if isinstance(paras, list) else [paras]):
        if not isinstance(para, P):
            para = P(para)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = para.align or align
        p.space_after = Pt(para.after)
        p.space_before = Pt(para.before)
        p.line_spacing = para.line
        if para.bullet:
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(int(Inches(para.hang))))
            pPr.set("indent", str(-int(Inches(para.hang))))
            clr = pPr.makeelement(qn("a:buClr"), {})
            clr.append(clr.makeelement(qn("a:srgbClr"), {"val": str(para.bcolor or BLUE)}))
            pPr.append(clr)
            pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
            pPr.append(pPr.makeelement(qn("a:buChar"), {"char": para.bullet}))
        for text, o in para.runs:
            r = p.add_run()
            r.text = text
            style_run(r, o.get("size", size), o.get("font"), o.get("color", color), o.get("italic", False))
    return tb


def card(slide, x, y, w, h, fill=WHITE, line=BORDER, radius=0.07, name=None, width=0.75):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        sh.name = name
    sh.adjustments[0] = min(0.5, radius / min(w, h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(width)
    sh.shadow.inherit = False
    return sh


def pill(slide, x, y, w, h, text, fill=TINT, color=BLUE, size=10, font="b", line=None):
    card(slide, x, y, w, h, fill=fill, line=line, radius=h / 2)
    tx(slide, x, y, w, h, [P([(text, {"size": size, "font": font, "color": color})], align=PP_ALIGN.CENTER)], anchor=MSO_ANCHOR.MIDDLE)


def pic(slide, path, x, y, w=None, h=None):
    iw, ih = Image.open(path).size
    if w is None:
        w = h * iw / ih
    if h is None:
        h = w * ih / iw
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))


def icon(slide, name, x, y, size=0.3, color="blue"):
    return pic(slide, IMG / f"i-{name}-{color}.png", x, y, w=size, h=size)


def icon_chip(slide, name, x, y, size=0.42, fill=TINT, color="blue"):
    card(slide, x, y, size, size, fill=fill, line=None, radius=size * 0.28)
    pad = size * 0.24
    icon(slide, name, x + pad, y + pad, size - 2 * pad, color)


def line(slide, x1, y1, x2, y2, color=MUTED, width=1.25, head=False, dash=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    st = c._element.find(qn("p:style"))
    if st is not None:  # the theme's effect style would add a drop shadow to every line
        st.find(qn("a:effectRef")).set("idx", "0")
    ln = c.line._get_or_add_ln()
    if dash:
        ln.append(ln.makeelement(qn("a:prstDash"), {"val": "dash"}))
    if head:
        ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "lg", "len": "med"}))
    return c


def polyline(slide, pts, color=MUTED, width=1.25, dash=False):
    for i in range(len(pts) - 1):
        line(slide, *pts[i], *pts[i + 1], color=color, width=width, head=(i == len(pts) - 2), dash=dash)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def title_style(slide):
    """Keep the organisers' title text and position; only swap the missing Google font for the deck font and use navy."""
    for sh in slide.shapes:
        if sh.name.startswith("Google Shape") and sh.has_text_frame and sh.text_frame.text.strip() and Inches(0.6) < sh.top < Inches(0.9):
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    r.font.name = FONT_BOLD
                    r.font.color.rgb = NAVY
            return sh


def slide_at(prs, n):
    s = prs.slides[n - 1]
    return s


# ---------- composed images ----------
def _shadowed(canvas_size, box, radius, bg, blur=11, offset=(0, 7), alpha=70):
    sh = Image.new("RGBA", canvas_size, (0, 0, 0, 0))
    d = ImageDraw.Draw(sh)
    x0, y0, x1, y1 = box
    d.rounded_rectangle((x0 + offset[0], y0 + offset[1], x1 + offset[0], y1 + offset[1]), radius, fill=(8, 27, 50, alpha))
    sh = sh.filter(ImageFilter.GaussianBlur(blur))
    base = Image.new("RGBA", canvas_size, bg + (255,))
    base.alpha_composite(sh)
    return base


def phone_image(name: str, src: str, bg=(255, 255, 255), crop=(0, 0, 780, 1688), screen_w=396) -> Path:
    """A screenshot inside a phone bezel with a soft shadow, flattened onto bg (JPEG keeps the PDF small)."""
    GEN.mkdir(parents=True, exist_ok=True)
    out = GEN / f"phone-{name}-{'%02x%02x%02x' % bg}.jpg"
    shot = Image.open(SHOTS / src).convert("RGB").crop(crop)
    sw = screen_w
    sh_ = round(shot.height * sw / shot.width)
    shot = shot.resize((sw, sh_), Image.LANCZOS)
    bez, pad = 9, 26
    W, H = sw + 2 * bez + 2 * pad, sh_ + 2 * bez + 2 * pad
    ro, ri = 40, 31
    canvas = _shadowed((W, H), (pad, pad, W - pad, H - pad), ro, bg)
    d = ImageDraw.Draw(canvas)
    d.rounded_rectangle((pad, pad, W - pad - 1, H - pad - 1), ro, fill=(11, 31, 51, 255))
    d.rounded_rectangle((pad + 1.5, pad + 1.5, W - pad - 2.5, H - pad - 2.5), ro - 1, outline=(70, 88, 112, 255), width=1)
    mask = Image.new("L", shot.size, 0)
    ImageDraw.Draw(mask).rounded_rectangle((0, 0, shot.width - 1, shot.height - 1), ri, fill=255)
    canvas.paste(shot.convert("RGBA"), (pad + bez, pad + bez), mask)
    canvas.convert("RGB").save(out, quality=88, optimize=True)
    return out


def desktop_image(name: str, src: str, bg=(255, 255, 255), crop=(0, 0, 1280, 800), w=1000) -> Path:
    """A screenshot inside a browser window frame."""
    GEN.mkdir(parents=True, exist_ok=True)
    out = GEN / f"desk-{name}-{'%02x%02x%02x' % bg}.jpg"
    shot = Image.open(SHOTS / src).convert("RGB").crop(crop)
    shot = shot.resize((w, round(shot.height * w / shot.width)), Image.LANCZOS)
    bar, pad = 34, 24
    W, H = w + 2 * pad, shot.height + bar + 2 * pad
    canvas = _shadowed((W, H), (pad, pad, W - pad, H - pad), 14, bg, blur=10, offset=(0, 6), alpha=60)
    d = ImageDraw.Draw(canvas)
    d.rounded_rectangle((pad, pad, W - pad - 1, H - pad - 1), 14, fill=(238, 242, 248, 255), outline=(200, 211, 226, 255))
    for i, c in enumerate([(236, 106, 94), (244, 191, 79), (97, 197, 84)]):
        cx = pad + 20 + i * 20
        d.ellipse((cx - 6, pad + bar / 2 - 6, cx + 6, pad + bar / 2 + 6), fill=c)
    d.rounded_rectangle((pad + 90, pad + 8, W - pad - 90, pad + bar - 8), 9, fill=(255, 255, 255, 255), outline=(214, 222, 234, 255))
    body = Image.new("RGBA", (w, shot.height), (255, 255, 255, 255))
    body.paste(shot, (0, 0))
    mask = Image.new("L", body.size, 255)
    md = ImageDraw.Draw(mask)
    md.rectangle((0, 0, w, shot.height), fill=255)
    canvas.paste(body, (pad, pad + bar))
    # re-round the bottom corners of the window
    corner = Image.new("L", (W, H), 0)
    ImageDraw.Draw(corner).rounded_rectangle((pad, pad, W - pad - 1, H - pad - 1), 14, fill=255)
    base = _shadowed((W, H), (pad, pad, W - pad, H - pad), 14, bg, blur=10, offset=(0, 6), alpha=60)
    base.paste(canvas, (0, 0), corner)
    base.convert("RGB").save(out, quality=88, optimize=True)
    return out


def qr_image(url: str) -> Path:
    import qrcode

    GEN.mkdir(parents=True, exist_ok=True)
    out = GEN / "qr-live.png"
    qr = qrcode.QRCode(error_correction=qrcode.constants.ERROR_CORRECT_M, box_size=12, border=2)
    qr.add_data(url)
    qr.make(fit=True)
    qr.make_image(fill_color="#081B32", back_color="white").convert("RGB").save(out, optimize=True)
    return out


def phone(slide, name, src, x, y, h, bg=(255, 255, 255), crop=(0, 0, 780, 1688)):
    path = phone_image(name, src, bg, crop)
    iw, ih = Image.open(path).size
    return pic(slide, path, x, y, w=h * iw / ih, h=h), h * iw / ih


# ---------- slides ----------
def slide1(prs):
    s = slide_at(prs, 1)
    box = next(sh for sh in s.shapes if sh.has_text_frame and "Team name" in sh.text_frame.text)
    box.left, box.top, box.width, box.height = Inches(0.16), Inches(3.08), Inches(9.68), Inches(2.42)
    answers = {"Team name:": "Scam Shield", "Team leader name:": "Naman Shrimal"}
    paras = box.text_frame.paragraphs
    for p in paras:
        t = p.text.strip()
        if t == "Team Details":
            for r in p.runs:
                r.font.name, r.font.size = FONT_BOLD, Pt(20)
                r.font.color.rgb = NAVY
        elif t.startswith(("Team name", "Team leader", "Problem Statement")):
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(int(Inches(0.42))))
            pPr.set("indent", str(-int(Inches(0.3))))
            for bu in pPr.findall(qn("a:buSzPts")):
                bu.set("val", "1300")
            for r in p.runs:
                style_run(r, 13, None, MUTED)
            p.space_after = Pt(3)
            label = next((k for k in answers if t.startswith(k)), None)
            if label:
                r = p.add_run()
                r.text = answers[label]
                style_run(r, 14, "b", INK)
            else:
                p.space_before = Pt(2)
                p.add_line_break()
                r = p.add_run()
                r.text = TRACK
                style_run(r, 13, "b", BLUE)
                p.add_line_break()
                r = p.add_run()
                r.text = ("People in India lose crores to scam messages, QR codes and calls that look real. Scam Shield lets anyone check "
                          "what they were sent in seconds, then act: a complaint draft, a shared scam list that warns the next person, and a "
                          "WhatsApp card for the family. Community benefit: fewer people lose their savings, and every report protects the next one.")
                style_run(r, 12, None, INK)
        elif not t:  # the template's two empty spacer paragraphs: keep them, but make them small
            for er in p._p.findall(qn("a:endParaRPr")):
                er.set("sz", "600")
            for r in p.runs:
                r.font.size = Pt(6)
    # drop the trailing spacer so the statement never runs into the footer line
    last = paras[-1]
    if not last.text.strip():
        last._p.getparent().remove(last._p)
    pic(s, IMG / "logo-lockup-horizontal.png", 7.1, 3.12, w=2.5)
    notes(s, "Team: Scam Shield. Team leader: Naman Shrimal. Track: BFSI: Intelligent Risk, Fraud & Financial Experiences. "
             "Community benefit: a free, plain-language check for anyone in India, in English or Hindi, and a shared scam list so each report warns the next person. "
             "The statement 'lose crores' is backed by slide 4 (about Rs 19,813 crore reported lost to cyber fraud in 2025, I4C data; see docs/submission/sources.md).")


def slide3(prs):  # Brief about the idea
    s = slide_at(prs, 3)
    title_style(s)
    lw = 5.55
    tx(s, MX, Y0, lw, 0.78, [
        P([("Check what you were just sent before you act on it.", {"size": 16, "font": "b", "color": NAVY})], after=3),
        P([("Free for anyone in India, and built for the scams people meet here: fake payment receipts, \"digital arrest\" calls and UPI QR tricks.", {"size": 12, "color": INK})]),
    ])
    items = [
        ("sparkles", "One multimodal Gemini call", "Text, screenshots, QR codes, links and Hinglish voice notes are read together by Gemini on Vertex AI."),
        ("quote", "A verdict you can check", "Scam, Suspicious or No scam signs found. Every red flag quotes the exact words. English or Hindi."),
        ("flag", "From verdict to action", "A complaint draft for 1930 and cybercrime.gov.in, a shared scam list and a WhatsApp warning card."),
    ]
    y, ch3 = 2.17, 0.8
    for ic, t, b in items:
        card(s, MX, y, lw, ch3)
        icon_chip(s, ic, MX + 0.16, y + (ch3 - 0.42) / 2)
        tx(s, MX + 0.78, y, lw - 0.92, ch3, [P([(t, {"size": 12.5, "font": "b", "color": NAVY})], after=1), P([(b, {"size": 11.5, "color": INK})])], anchor=MSO_ANCHOR.MIDDLE)
        y += ch3 + 0.07
    card(s, MX, y, lw, Y1 - y, fill=NAVY, line=None, radius=0.08)
    tx(s, MX + 0.2, y, lw - 0.4, Y1 - y, [P([("Careful by design: ", {"size": 12, "font": "b", "color": CYAN}),
                                           ("links open in an isolated browser, never on your phone. It never says \"safe\": it gives guidance and you decide.", {"size": 12, "color": WHITE})])], anchor=MSO_ANCHOR.MIDDLE)
    px = 6.25
    card(s, px - 0.1, Y0, 3.45, 3.62, fill=TINT2, line=None, radius=0.12)
    bg = (0xF1, 0xF5, 0xFB)
    _, w1 = phone(s, "landing", "02a-landing-mobile.png", px, Y0 + 0.13, 3.38, bg)
    phone(s, "scam", "03-result-scam-mobile.png", px + w1 - 0.05, Y0 + 0.13, 3.38, bg)
    tx(s, px - 0.1, 5.02, 3.45, 0.3, [P([("The live app, as seen on a phone.", {"size": 11.5, "color": MUTED})], align=PP_ALIGN.CENTER)], anchor=MSO_ANCHOR.MIDDLE)
    notes(s, "Gen AI (the heart of the idea): one multimodal call to Gemini on Vertex AI reads the text, a screenshot or an audio clip together (Hinglish voice notes are transcribed and judged in the same call). "
             "The answer is schema-validated structured output, every red flag quotes evidence from the input, and the answer comes back in English or Hindi. "
             "Deterministic checks (scam-list hits, a QR that asks you to pay while the text says receive) can only RAISE a verdict. "
             "What Gemini does that rules cannot: judge scams nobody has reported yet, and read images and voice. "
             "India-specific scam types covered by the benchmark and the examples: fake payment receipts, 'digital arrest' calls, fake KYC messages, UPI QR tricks. "
             "Careful by design: links are opened only in an isolated scanner (a private Cloud Run service), never on the person's phone; verdicts are Scam, Suspicious or No scam signs found, the app never says 'safe', and it gives guidance only. "
             "Screenshots are from the live production app, captured by scripts/screenshots.ts.")


def slide4(prs):  # Opportunities
    s = slide_at(prs, 4)
    t = title_style(s)
    # The organisers' sub-prompts were lost with the original file; the three cards below answer them under their standard wording.
    for p in t.text_frame.paragraphs:
        for br in p._p.findall(qn("a:br")):
            p._p.remove(br)
    card(s, MX, Y0, CW, 0.58, fill=NAVY, line=None)
    tx(s, MX + 0.22, Y0, 2.75, 0.58, [P([("Rs 19,813 crore", {"size": 21, "font": "b", "color": CYAN})])], anchor=MSO_ANCHOR.MIDDLE)
    tx(s, MX + 3.05, Y0, CW - 3.25, 0.58, [P([("reported lost to cyber fraud in 2025, across about 21.78 lakh complaints (I4C data, National Cyber Crime Reporting Portal)", {"size": 12, "color": WHITE})])], anchor=MSO_ANCHOR.MIDDLE)
    top = 2.0
    lw = 5.35
    card(s, MX, top, lw, Y1 - top)
    tx(s, MX + 0.2, top + 0.14, lw - 0.4, 0.26, [P([("How different is it from existing ideas?", {"size": 13, "font": "b", "color": BLUE})])])
    rows = [
        ("Truecaller Scam Checker (Sep 2026): ", "checks links, numbers and messages against community reports; screenshot upload is only planned.", 0.64),
        ("Google scam detection: ", "Messages on Pixel 6 and newer, and calls on Pixel 9 and newer, in India.", 0.44),
        ("Chakshu (Sanchar Saathi): ", "a reporting facility for suspicious calls, SMS and WhatsApp messages.", 0.44),
    ]
    y = top + 0.5
    for name, body, h in rows:
        tx(s, MX + 0.2, y, lw - 0.4, h, [P([(name, {"size": 12, "font": "b", "color": NAVY}), (body, {"size": 12, "color": INK})])])
        y += h + 0.04
        line(s, MX + 0.2, y, MX + lw - 0.2, y, color=BORDER, width=0.75)
        y += 0.08
    card(s, MX + 0.14, y + 0.02, lw - 0.28, Y1 - 0.14 - y - 0.02, fill=TINT, line=None, radius=0.06)
    tx(s, MX + 0.28, y + 0.02, lw - 0.56, Y1 - 0.14 - y - 0.02, [P([("Scam Shield: ", {"size": 12, "font": "b", "color": BLUE}), ("reads the content itself, so it can flag scams nobody has reported yet, from text, screenshots, QR codes, links and voice notes.", {"size": 12, "color": INK})])], anchor=MSO_ANCHOR.MIDDLE)
    rx, rw = MX + lw + 0.2, CW - lw - 0.2
    ha = 1.44
    card(s, rx, top, rw, ha)
    tx(s, rx + 0.2, top + 0.14, rw - 0.4, 0.26, [P([("How will it solve the problem?", {"size": 13, "font": "b", "color": BLUE})])])
    tx(s, rx + 0.2, top + 0.48, rw - 0.4, ha - 0.55, [
        P("Checks before you pay, scan or click", bullet="•", after=3),
        P("Quotes every red flag from the message", bullet="•", after=3),
        P("Turns a scam into action: complaint, shared list, family warning", bullet="•"),
    ], size=12)
    y2 = top + ha + 0.12
    card(s, rx, y2, rw, Y1 - y2)
    tx(s, rx + 0.2, y2 + 0.14, rw - 0.4, 0.26, [P([("USP of the proposed solution", {"size": 13, "font": "b", "color": BLUE})])])
    tx(s, rx + 0.2, y2 + 0.48, rw - 0.4, Y1 - y2 - 0.55, [
        P("Catches scams nobody has reported yet", bullet="•", after=3),
        P("Made for India: fake payment receipts, \"digital arrest\" calls, UPI QR tricks", bullet="•", after=3),
        P("Understands Hinglish voice notes", bullet="•"),
    ], size=12)
    notes(s, "Sources (all checked 18 Oct 2026, see docs/submission/sources.md): "
             "Rs 19,813 crore and about 21.78 lakh complaints reported for 2025, I4C / National Cyber Crime Reporting Portal (https://ianslive.in/indians-lose-over-rs-52976-crore-to-cyber-frauds-over-six-years-report--20260103154943; a separate Home Ministry reply is quoted as about Rs 22,495 crore, so the figure is labelled I4C data). "
             "Truecaller Scam Checker launched in India on 28 Sep 2026: checks links, numbers and messages against community reports; screenshot upload is planned, not available (https://www.techrepublic.com/article/news-truecaller-scam-checker-apac-india/). "
             "Google scam detection: Messages on Pixel 6 and newer reached India in Sep 2026; on-device call scam detection in India is limited to Pixel 9 and newer (https://androidauthority.com/pixel-scam-detection-google-messages-global-3711361, https://techcrunch.com/2025/11/20/google-steps-up-ai-scam-protection-in-india-but-gaps-remain). "
             "Chakshu (Sanchar Saathi) is a reporting facility (https://sancharsaathi.gov.in; wording to be re-verified). "
             "No recovery percentage is quoted: no national recovery rate is published. "
             "USP wording: fake payment receipts and 'digital arrest' calls are scam types in the benchmark and the examples; UPI QR pay-vs-receive is a hard check. "
             "The three card headings are the organisers' standard prompts for this slide (the original prompt text was lost from the template file).")


def slide5(prs):  # Features
    s = slide_at(prs, 5)
    title_style(s)
    feats = [
        ("scan-text", "5 input types", "Text, screenshot, QR code, link and voice note."),
        ("shield-check", "Clear verdict", "Scam, Suspicious or No scam signs found, with a risk score."),
        ("quote", "Evidence quoted", "Each red flag points to the exact words in the message."),
        ("languages", "English + Hindi", "Hinglish understood; answers in the language you pick."),
        ("file-text", "Complaint draft", "Ready text for 1930 and cybercrime.gov.in."),
        ("users", "Shared scam list", "\"Reported N times\", with linked numbers, UPI IDs and sites."),
        ("share-2", "Warning card", "One tap shares a warning with the family WhatsApp group."),
        ("bot", "Agent tools", "WebMCP tools let AI agents check and report."),
    ]
    cw, ch, gx, gy = 2.15, 1.72, 0.2, 0.12
    for i, (ic, t, b) in enumerate(feats):
        col, row = i % 4, i // 4
        x, y = MX + col * (cw + gx), Y0 + row * (ch + gy)
        card(s, x, y, cw, ch)
        icon_chip(s, ic, x + 0.16, y + 0.16, 0.4)
        tx(s, x + 0.16, y + 0.68, cw - 0.3, ch - 0.74, [P([(t, {"size": 13, "font": "b", "color": NAVY})], after=2), P([(b, {"size": 12, "color": INK})])])
    yb = Y0 + 2 * ch + gy + 0.1
    card(s, MX, yb, CW, Y1 - yb, fill=TINT2, line=None, radius=0.07)
    tx(s, MX + 0.2, yb, CW - 0.4, Y1 - yb, [P([("Also: ", {"size": 12, "font": "b", "color": BLUE}), ("trends, history, one-tap examples, an installable app and the Android share menu.", {"size": 12, "color": INK})])], anchor=MSO_ANCHOR.MIDDLE)
    notes(s, "All features are live in the production app (https://project-91260b43-6705-4176-b71.web.app). The WebMCP tools check_message and report_scam are registered in web/src/webmcp/registerTools.ts; "
             "report_scam needs a person signed in with Google. Verdict names are Scam, Suspicious and No scam signs found.")


def slide6(prs):  # Process flow
    s = slide_at(prs, 6)
    title_style(s)
    steps = [
        ("Input", "Text, screenshot, QR, link or voice note"),
        ("Guardrails", "Anonymous sign-in; per-user, per-IP and daily limits"),
        ("Extract", "Pulls out phones, UPI IDs, links; blanks OTPs and cards"),
        ("Scam-list lookup", "Firestore: \"Reported N times\" and linked IDs"),
        ("Isolated link scan", "A private Cloud Run service opens it: URL, text, screenshot"),
        ("Gemini verdict", "One call reads text, image or audio; evidence quoted"),
        ("Hard checks", "Scam-list hits and QR pay-vs-receive only raise the verdict"),
        ("Result and action", "Verdict, complaint draft, shared list, warning card"),
    ]
    bw, bh, gap = 2.15, 1.4, 0.2
    ys = [Y0, Y0 + bh + 0.22]
    for i, (t, b) in enumerate(steps):
        col, row = i % 4, i // 4
        x, y = MX + col * (bw + gap), ys[row]
        gem = i == 5
        card(s, x, y, bw, bh, fill=NAVY if gem else WHITE, line=None if gem else BORDER)
        card(s, x + 0.15, y + 0.14, 0.28, 0.28, fill=CYAN if gem else BLUE, line=None, radius=0.14)
        tx(s, x + 0.15, y + 0.14, 0.28, 0.28, [P([(str(i + 1), {"size": 11, "font": "b", "color": NAVY if gem else WHITE})], align=PP_ALIGN.CENTER)], anchor=MSO_ANCHOR.MIDDLE)
        if gem:
            pill(s, x + bw - 0.78, y + 0.15, 0.63, 0.26, "Gen AI", fill=CYAN, color=NAVY, size=10)
        tx(s, x + 0.15, y + 0.47, bw - 0.28, bh - 0.52, [P([(t, {"size": 12.5, "font": "b", "color": WHITE if gem else NAVY})], after=2), P([(b, {"size": 11.5, "color": WHITE if gem else INK})])])
        if col < 3:
            line(s, x + bw + 0.03, y + bh / 2, x + bw + gap - 0.03, y + bh / 2, color=BLUE, width=1.75, head=True)
    x4, x5 = MX + 3 * (bw + gap) + bw / 2, MX + bw / 2
    ym = Y0 + bh + 0.11
    line(s, x4, Y0 + bh, x4, ym, color=BLUE, width=1.75)
    line(s, x4, ym, x5, ym, color=BLUE, width=1.75)
    line(s, x5, ym, x5, ys[1], color=BLUE, width=1.75, head=True)
    yb = ys[1] + bh + 0.12
    card(s, MX, yb, CW, Y1 - yb, fill=TINT, line=None, radius=0.07)
    icon(s, "brain", MX + 0.2, yb + (Y1 - yb) / 2 - 0.17, 0.34)
    tx(s, MX + 0.7, yb, CW - 0.9, Y1 - yb, [P([("Why Gen AI: ", {"size": 11.5, "font": "b", "color": BLUE}), ("rules cannot read a screenshot or a Hinglish voice note, or judge a scam nobody has reported. Gemini does. The hard checks can only raise its verdict, never lower it.", {"size": 11.5, "color": INK})], after=3),
                                         P([("If Gemini fails or times out, the person sees a clear retry message; no verdict is invented and nothing is stored.", {"size": 11.5, "color": INK})])], anchor=MSO_ANCHOR.MIDDLE)
    notes(s, "Flow as implemented in api/src/app.ts. Step 6 is a single multimodal Gemini call on Vertex AI (default gemini-3.7-flash, fallback gemini-2.5-flash) that returns schema-validated structured output with quoted evidence in English or Hindi. "
             "Step 7: deterministic checks (scam-list hits, a QR that asks you to pay while the text says receive) can only RAISE the verdict. "
             "If Gemini fails or times out the person sees a clear retry message; no verdict is invented and nothing is stored. "
             "Limits: 30 checks per user per day, 60 per IP per minute, 3,000 per day overall.")


def slide7(prs):  # Wireframes
    s = slide_at(prs, 7)
    title_style(s)
    bg = (255, 255, 255)
    my, mh = 1.3, 3.4
    dpath = desktop_image("check", "01-check-desktop.png", bg)
    iw, ih = Image.open(dpath).size
    dw = 3.95
    dh = dw * ih / iw
    gap = 0.2
    pws = []
    items = [("check", "02-check-mobile.png"), ("complaint", "06-complaint-mobile.png"), ("hindi", "09-check-hindi-mobile.png")]
    for n, f in items:
        pi = phone_image(n, f, bg)
        pw_, ph_ = Image.open(pi).size
        pws.append(mh * pw_ / ph_)
    total = dw + sum(pws) + gap * 3
    x = MX + (CW - total) / 2
    pic(s, dpath, x, my + (mh - dh) / 2, w=dw)
    caps = [("Web app", "The same flow on desktop, with a sidebar.")]
    xs = [x]
    x += dw + gap
    for (n, f), w_ in zip(items, pws):
        phone(s, n, f, x, my, mh, bg)
        xs.append(x)
        x += w_ + gap
    ws = [dw] + pws
    caps += [("1  Paste or share", "Text, photo, QR or voice note."), ("2  Act on a scam", "A complaint draft ready to file."), ("3  Hindi", "The whole app switches language.")]
    for (t, b), cx, cw_ in zip(caps, xs, ws):
        tx(s, cx + 0.1, 4.78, cw_ - 0.2, 0.6, [P([(t, {"size": 12, "font": "b", "color": NAVY})], after=1), P([(b, {"size": 11, "color": MUTED})])], align=PP_ALIGN.CENTER)
    notes(s, "These are real screenshots of the running app, not mock-ups: the desktop check page, the mobile check page, the complaint page and the Hindi interface. "
             "Captured with scripts/screenshots.ts against the live production site.")


def node(slide, x, y, w, h, title, sub, logos=(), fill=WHITE, line_=BORDER, tcolor=NAVY, scolor=MUTED, tsize=12, ssize=11, side=True):
    """A diagram box. Logos sit left of the text (side=True) or above it."""
    card(slide, x, y, w, h, fill=fill, line=line_ if fill != NAVY else None)
    lx = x + 0.14
    lh = 0.3
    if side and logos:
        for lg in logos:
            p_ = pic(slide, IMG / lg, lx, y + 0.14, h=lh)
            lx += p_.width / 914400 + 0.06
        tx0, ty = lx + 0.08, y + 0.13
    elif logos:
        for lg in logos:
            p_ = pic(slide, IMG / lg, lx, y + 0.13, h=lh)
            lx += p_.width / 914400 + 0.08
        tx0, ty = x + 0.14, y + 0.52
    else:
        tx0, ty = x + 0.14, y + 0.12
    tx(slide, tx0, ty, x + w - tx0 - 0.1, y + h - ty - 0.04, [P([(title, {"size": tsize, "font": "b", "color": tcolor})], after=1), P([(sub, {"size": ssize, "color": scolor})])])


YF = 5.14  # content bottom on slides 8 and 9, which carry the trademark footnote between the content and the footer bar


def tm_footnote(slide):
    """Trademark and non-affiliation footnote for slides that show Google product icons (9 pt, clear of the footer bar)."""
    tx(slide, MX, 5.22, CW, 0.22, [P([("Product names and icons are trademarks of Google LLC, shown only to identify the technologies used. No endorsement or affiliation implied.", {"size": 9, "color": MUTED})])], anchor=MSO_ANCHOR.MIDDLE)


def slide8(prs):  # Architecture
    s = slide_at(prs, 8)
    title_style(s)
    ax, aw = MX, 1.8
    bx, bw = 2.6, 2.5
    cx, cw = 5.6, 4.0
    node(s, ax, 1.35, aw, 2.1, "Person", "Browser or installed app (PWA), English and Hindi, Android share menu. Uploads go straight to Storage.", ("i-smartphone-blue.png",), side=False, ssize=11.5)
    node(s, ax, 3.7, aw, YF - 3.7, "Firebase Auth", "Anonymous and Google sign-in", ("g-firebase.png",), side=False, ssize=11.5)
    node(s, bx, 1.35, bw, 0.95, "Firebase Hosting", "Web app; /api/** goes to the API", ("g-firebase.png",))
    node(s, bx, 2.55, bw, 1.4, "Cloud Run: API", "Stateless; 1 to 5 instances, one always warm. Limits, extraction, hard checks", ("g-cloud-run.png",), fill=NAVY, tcolor=WHITE, scolor=rgb("#C9D6EA"))
    node(s, bx, 4.15, bw, YF - 4.15, "Cost guard", "30 checks per user per day, 60 per IP per minute, 3,000 per day", (), fill=TINT, line_=None)
    node(s, cx, 1.35, cw, 0.9, "Gemini on Vertex AI", "One multimodal call: text, image or audio in, structured verdict out", ("g-gemini.png", "g-vertex-ai.png"))
    node(s, cx, 2.35, cw, 0.78, "Firestore", "Shared scam list, checks, reports, trend counters", ("g-firestore.png",))
    node(s, cx, 3.23, cw, 0.88, "Cloud Run: link scanner (private)", "Isolated Chromium; scales to zero; zero-permission service account", ("g-cloud-run.png",))
    node(s, cx, 4.36, 1.95, YF - 4.36, "Cloud Storage", "Uploads, deleted after 7 days", (), tsize=11.5, ssize=10.5)
    node(s, cx + 2.05, 4.36, 1.95, YF - 4.36, "Suspicious site", "Opened only in the scanner", (), fill=RED_TINT, line_=None, tcolor=RED, scolor=INK, tsize=11.5, ssize=10.5)
    # arrows
    line(s, ax + aw, 1.82, bx, 1.82, color=BLUE, width=1.75, head=True)
    line(s, bx + bw / 2, 2.3, bx + bw / 2, 2.55, color=BLUE, width=1.75, head=True)
    line(s, ax + aw / 2, 3.45, ax + aw / 2, 3.7, color=BLUE, width=1.75, head=True)
    polyline(s, [(bx + bw, 2.7), (5.35, 2.7), (5.35, 1.8), (cx, 1.8)], color=BLUE, width=1.75)
    line(s, bx + bw, 2.95, cx, 2.95, color=BLUE, width=1.75, head=True)
    line(s, bx + bw, 3.5, cx, 3.5, color=BLUE, width=1.75, head=True)
    polyline(s, [(bx + bw, 3.78), (5.35, 3.78), (5.35, 4.75), (cx, 4.75)], color=BLUE, width=1.75)
    line(s, cx + 3.0, 4.11, cx + 3.0, 4.36, color=RED, width=1.75, head=True)
    tm_footnote(s)
    notes(s, "Architecture as deployed (scripts/deploy.sh, region asia-south1): Firebase Hosting serves the React PWA and forwards /api/** to the Cloud Run API. "
             "The API is stateless (1 vCPU, 1 GiB, concurrency 20, min 1 and max 5 instances) and calls Gemini on Vertex AI, Firestore and the private link scanner. "
             "The link scanner (2 vCPU, 2 GiB, min 0, max 3 instances, private, no unauthenticated access) runs under its own service account with no project roles and can be invoked only by the API's service account. "
             "Cloud Storage uploads are deleted after 7 days by a lifecycle rule (scripts/storage-lifecycle.json). Limits: 30 per user per day, 60 per IP per minute, 3,000 per day overall. "
             "Product icons are shown only to identify the technologies used; no endorsement by Google is implied (trademark footnote on the slide). The icon files are kept in docs/brand/third-party. Uploads go straight from the browser to Storage.")


def chips(slide, x, y, w, labels, size=11, h=0.28, gap=0.08, fill=TINT2, color=INK):
    cx, cy = x, y
    for lab in labels:
        cwid = 0.2 + 0.076 * len(lab) * size / 11
        if cx + cwid > x + w:
            cx, cy = x, cy + h + gap
        pill(slide, cx, cy, cwid, h, lab, fill=fill, color=color, size=size, font="m", line=BORDER)
        cx += cwid + gap
    return cy + h


def slide9(prs):  # Technologies
    s = slide_at(prs, 9)
    title_style(s)
    lw = 4.3
    card(s, MX, Y0, lw, YF - Y0, fill=NAVY, line=None, radius=0.09)
    # white tile behind the colourful marks so they stay readable on navy
    card(s, MX + 0.2, Y0 + 0.2, 1.12, 0.46, fill=WHITE, line=None, radius=0.1)
    pic(s, IMG / "g-gemini.png", MX + 0.3, Y0 + 0.28, h=0.3)
    pic(s, IMG / "g-vertex-ai.png", MX + 0.76, Y0 + 0.28, h=0.3)
    tx(s, MX + 1.45, Y0 + 0.2, lw - 1.6, 0.46, [P([("Gen AI that does real work", {"size": 14, "font": "b", "color": WHITE})])], anchor=MSO_ANCHOR.MIDDLE)
    bl = [
        ("Gemini 3.7 Flash ", "on Vertex AI, with Gemini 2.5 Flash as automatic fallback"),
        ("One multimodal call ", "reads text, screenshots and audio (Hinglish voice)"),
        ("Schema-validated output; ", "every red flag quotes evidence; English or Hindi"),
        ("Hard checks only raise ", "the verdict: scam list, QR pay-vs-receive"),
        ("Does what rules cannot: ", "unreported scams, images, voice"),
        ("WebMCP agent tools: ", "check_message and report_scam"),
    ]
    tx(s, MX + 0.22, Y0 + 0.88, lw - 0.44, YF - Y0 - 1.0, [P([(a, {"size": 12, "font": "b", "color": CYAN}), (b, {"size": 12, "color": WHITE})], bullet="•", bcolor=CYAN, after=6) for a, b in bl])
    rx, rw = MX + lw + 0.2, CW - lw - 0.2
    ht = 2.18
    card(s, rx, Y0, rw, ht)
    tx(s, rx + 0.2, Y0 + 0.14, rw - 0.4, 0.26, [P([("Google Cloud and Firebase, built to scale", {"size": 13, "font": "b", "color": BLUE})])])
    lx = rx + 0.2
    for lg in ("g-google-cloud.png", "g-cloud-run.png", "g-firestore.png", "g-firebase.png"):
        p_ = pic(s, IMG / lg, lx, Y0 + 0.5, h=0.3)
        lx += p_.width / 914400 + 0.2
    tx(s, rx + 0.2, Y0 + 0.9, rw - 0.4, ht - 0.95, [
        P([("Cloud Run API: ", {"size": 12, "font": "b", "color": NAVY}), ("stateless and autoscaling", {"size": 12})], bullet="•", after=4),
        P([("Link scanner: ", {"size": 12, "font": "b", "color": NAVY}), ("scales to zero, own zero-permission service account", {"size": 12})], bullet="•", after=4),
        P([("Firestore ", {"size": 12, "font": "b", "color": NAVY}), ("shared scam list; ", {"size": 12}), ("Firebase ", {"size": 12, "font": "b", "color": NAVY}), ("Hosting, Auth, Storage", {"size": 12})], bullet="•"),
    ])
    y2 = Y0 + ht + 0.12
    card(s, rx, y2, rw, YF - y2)
    tx(s, rx + 0.2, y2 + 0.14, rw - 0.4, 0.26, [P([("Application", {"size": 13, "font": "b", "color": BLUE})])])
    cy = chips(s, rx + 0.2, y2 + 0.5, rw - 0.4, ["React 18", "Vite", "Tailwind", "Workbox PWA", "Node 22", "TypeScript", "Hono", "zod", "Playwright", "sharp + jsQR"], size=10.5, h=0.26, gap=0.07)
    tx(s, rx + 0.2, cy + 0.1, rw - 0.4, 0.3, [P([("590+ automated tests", {"size": 12, "font": "b", "color": NAVY}), (", plus a 50-case live benchmark", {"size": 12})])])
    tm_footnote(s)
    notes(s, "Meaningful Gen AI: a single multimodal Gemini call on Vertex AI (api/src/gemini.ts; default model gemini-3.7-flash, automatic fallback gemini-2.5-flash, location global) reads text, image or audio. "
             "Output is schema-validated (zod schemas in shared/), every red flag quotes evidence, answers come in English or Hindi. Deterministic checks can only raise the verdict. "
             "WebMCP tools: check_message and report_scam (web/src/webmcp/registerTools.ts). "
             "Scalability: the API is a stateless Cloud Run service (concurrency 20, min 1, max 5 instances); the link scanner scales to zero (max 3) under its own service account with no roles; Firestore holds the shared scam list; Firebase Hosting, Auth and Storage in front. "
             "Test count: on 18 Oct 2026 the suites passed 595 tests (shared 14, web 226, linkscan 94, eval 5, api 256); a further 18 api tests need the Firestore emulator (Java) and were not run. Also used: Cloud Build and Artifact Registry for deploys, Cloud Text-to-Speech to generate the voice example.")


def big_number_card(s, x, y, w, h, tag, number, unit, body, tag_fill=TINT, tag_color=BLUE):
    card(s, x, y, w, h)
    pill(s, x + 0.2, y + 0.16, 1.0 if len(tag) < 9 else 1.18, 0.26, tag, fill=tag_fill, color=tag_color, size=10)
    tx(s, x + 0.2, y + 0.5, w - 0.4, 0.45, [P([(number, {"size": 22, "font": "b", "color": NAVY})])])
    tx(s, x + 0.2, y + 0.95, w - 0.4, 0.25, [P([(unit, {"size": 11.5, "color": MUTED})])])
    line(s, x + 0.2, y + 1.26, x + w - 0.2, y + 1.26, color=BORDER, width=0.75)
    tx(s, x + 0.2, y + 1.34, w - 0.4, h - 1.4, [P([(body, {"size": 12, "color": INK})])])


def slide10(prs):  # Estimated cost
    s = slide_at(prs, 10)
    title_style(s)
    cw, gap, ch = 2.93, 0.2, 2.5
    big_number_card(s, MX, Y0, cw, ch, "FIXED", "about US$13", "per month (estimate)",
                    "One Cloud Run API instance kept warm (1 vCPU, 1 GiB, Mumbai) at US$0.0000025 per vCPU-second and per GiB-second, before any free tier.")
    big_number_card(s, MX + cw + gap, Y0, cw, ch, "PER CHECK", "US$0.004–0.008", "per text check (estimate)",
                    "Gemini 3.7 Flash at about 2,500 input and 500 to 1,500 output tokens. US$0.75 in, US$3.75 out per 1M tokens to 31 Dec 2026; double after.")
    big_number_card(s, MX + 2 * (cw + gap), Y0, cw, ch, "FREE TIERS", "Free tiers", "expected at this volume (estimate)",
                    "Firebase Hosting, Auth, Firestore and Storage are expected to stay inside free tiers. The link scanner bills only while it scans.")
    yb = Y0 + ch + 0.12
    card(s, MX, yb, CW, Y1 - yb, fill=NAVY, line=None, radius=0.08)
    cols = [
        ("Limits are the real control", "30 checks per user per day, 60 per IP per minute, 3,000 per day overall. These cap spending."),
        ("Ceiling at the daily cap", "3,000 checks at the per-check estimate is about US$11 to 23 per day for Gemini."),
        ("Budget alert: Rs 2,000 a month", "Alerts at 50%, 90% and 100% (Rs\u00a01,000, Rs\u00a01,800, Rs\u00a02,000). An alert notifies; it does not stop spending."),
    ]
    w3 = (CW - 0.4) / 3
    for i, (t, b) in enumerate(cols):
        x = MX + 0.2 + i * w3
        tx(s, x, yb + 0.2, w3 - 0.2, Y1 - yb - 0.25, [P([(t, {"size": 12.5, "font": "b", "color": CYAN})], after=3), P([(b, {"size": 12, "color": WHITE})])])
    notes(s, "Every figure on this slide is an ESTIMATE with stated assumptions, read from Google's published pages on 18 Oct 2026; none is a quote. "
             "Gemini: Vertex AI pricing page https://cloud.google.com/vertex-ai/generative-ai/pricing lists Gemini 3.7 Flash (global) at US$0.75 per 1M input tokens and US$3.75 per 1M output tokens through 31 Dec 2026, and US$1.50 / US$7.50 from 1 Jan 2027 (input includes text, image, video and audio; output includes reasoning). "
             "Assumption (ours, not measured): about 2,500 input and 500 output tokens per text check, up to 1,500 output tokens if reasoning tokens are counted. "
             "Low = 2,500 x 0.75/1M + 500 x 3.75/1M = US$0.0038; high = 2,500 x 0.75/1M + 1,500 x 3.75/1M = US$0.0075. Image and audio checks use more input tokens. At the 3,000 checks per day cap that is about US$11 to 23 per day. "
             "Cloud Run: https://cloud.google.com/run/pricing , asia-south1 (Mumbai) is Tier 1; idle minimum-instance time is US$0.0000025 per vCPU-second and per GiB-second (request-based billing). "
             "1 vCPU + 1 GiB x 2,592,000 s (30 days) x US$0.0000025 x 2 = US$12.96 per month before any free tier; active request time is billed on top and is small. "
             "Deployed API settings (scripts/deploy.sh): 1 vCPU, 1 GiB, min 1, max 5 instances. Free tier claims for Firebase Hosting, Auth, Firestore and Storage are expectations at this volume, not measured bills. "
             "Budget alert: Rs 2,000 per month, alerts at 50%, 90% and 100% (Rs 1,000, Rs 1,800, Rs 2,000), set by the owner in Google Cloud Billing (budget name 'Scam Shield', applies to this project). "
             "A budget alert only notifies; it does not cap or stop spending. The real cost control is the per-user (30 per day), per-IP (60 per minute) and global (3,000 per day) limits in the API. Currency: USD as listed by Google; no INR conversion is claimed.")


def slide11(prs):  # Snapshots
    s = slide_at(prs, 11)
    title_style(s)
    bgc = (0xF1, 0xF5, 0xFB)
    card(s, MX, Y0, CW, 3.38, fill=TINT2, line=None, radius=0.12)
    shots = [
        ("scam", "03-result-scam-mobile.png", "Scam", "Fake KYC SMS, evidence quoted", RED),
        ("qr", "04-result-qr-mobile.png", "Scam", "QR asks you to pay, text says receive", RED),
        ("nosigns", "05-result-no-signs-mobile.png", "No scam signs found", "A genuine bank alert; not a guarantee", SLATE),
        ("hivoice", "10-result-hindi-voice-mobile.png", "Scam, in Hindi", "Hinglish voice note, fake police call", RED),
        ("report", "07-report-dialog-mobile.png", "Report", "Adds it to the shared scam list", BLUE),
    ]
    ph = 3.22
    pw_ = [Image.open(phone_image(n, f, bgc)).size for n, f, *_ in shots]
    widths = [ph * w / h for w, h in pw_]
    gap = (CW - sum(widths)) / 6
    x = MX + gap
    for (n, f, t, b, col), w_ in zip(shots, widths):
        phone(s, n, f, x, Y0 + 0.08, ph, bgc)
        x += w_ + gap
    # captions below the panel, one per phone
    x = MX + gap
    for (n, f, t, b, col), w_ in zip(shots, widths):
        cap_w = w_ + gap - 0.08
        cx = x + w_ / 2 - cap_w / 2
        tx(s, cx, Y0 + 3.5, cap_w, 0.55, [P([(t, {"size": 12, "font": "b", "color": col if col != BLUE else NAVY})], after=1, align=PP_ALIGN.CENTER), P([(b, {"size": 11, "color": MUTED})], align=PP_ALIGN.CENTER)])
        x += w_ + gap
    notes(s, "Real screenshots of the production app, captured by scripts/screenshots.ts from live checks (the example buttons run the same API as a typed message). "
             "Left to right: a fake KYC SMS with a look-alike bank link; a QR that encodes a payment request while the text says 'receive'; a genuine bank debit alert with 'No scam signs found' (not a guarantee); "
             "a Hinglish 'digital arrest' voice note judged in Hindi; the report dialog, which asks for Google sign-in before a report is added to the shared list (to stop fake reports).")


def slide12(prs, res, hi):  # Performance
    s = slide_at(prs, 12)
    title_style(s)
    n, ns, ng = res["cases"]
    tiles = [("Accuracy", res["accuracy"]), ("Scam recall", res["recall"]), ("False-positive rate", res["fpr"]), ("Scam type correct", res["type"])]
    tw, gap, th = 2.15, 0.2, 0.84
    for i, (lab, val) in enumerate(tiles):
        x = MX + i * (tw + gap)
        card(s, x, Y0, tw, th)
        tx(s, x, Y0 + 0.07, tw, 0.42, [P([(val, {"size": 24, "font": "b", "color": BLUE})], align=PP_ALIGN.CENTER)])
        tx(s, x, Y0 + 0.5, tw, 0.26, [P([(lab, {"size": 11.5, "color": MUTED})], align=PP_ALIGN.CENTER)])
    tx(s, MX, Y0 + th + 0.07, CW, 0.22, [P([(f"English run, {res['date']}, {n} synthetic test cases, live app", {"size": 11.5, "color": MUTED})])])
    # latency table
    head = ["Input", "Cases", "English p50", "English p95"] + (["Hindi p50", "Hindi p95"] if hi else [])
    rows = [head]
    for k in ("text", "image", "audio"):
        if k in res["latency"]:
            p50, p95, cnt = res["latency"][k]
            row = [k.capitalize(), cnt, p50, p95]
            if hi:
                h50, h95, _ = hi["latency"].get(k, ("-", "-", ""))
                row += [h50, h95]
            rows.append(row)
    rh = 0.3
    tx_, ty_ = MX, 2.52
    lw_ = 5.55
    nr, nc = len(rows), len(head)
    tbl_shape = s.shapes.add_table(nr, nc, Inches(tx_), Inches(ty_), Inches(lw_), Inches(rh * nr))
    tbl = tbl_shape.table
    tblPr = tbl._tbl.tblPr
    for a in ("firstRow", "bandRow"):
        tblPr.set(a, "0")
    widths = [0.8, 0.6] + [(lw_ - 1.4) / (nc - 2)] * (nc - 2)
    for i, wd in enumerate(widths):
        tbl.columns[i].width = Inches(wd)
    for r in range(nr):
        tbl.rows[r].height = Inches(rh)
        for c in range(nc):
            cell = tbl.cell(r, c)
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
            run = p.add_run()
            run.text = str(rows[r][c])
            style_run(run, 11.5, "b" if (r == 0 or c == 0) else None, WHITE if r == 0 else INK)
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if r == 0 else (TINT2 if r % 2 else WHITE)
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    yb = ty_ + rh * nr + 0.1
    if hi:
        card(s, MX, yb, lw_, 0.5, fill=TINT, line=None, radius=0.07)
        tx(s, MX + 0.16, yb, lw_ - 0.3, 0.5, [P([("Hindi run: ", {"size": 11.5, "font": "b", "color": BLUE}),
                                                (f"accuracy {hi['accuracy']}, scam recall {hi['recall']}, false positives {hi['fpr']}, scam type correct {hi['type']}.", {"size": 11.5})])], anchor=MSO_ANCHOR.MIDDLE)
        yb += 0.6
    tx(s, MX, yb, lw_, Y1 - yb, [
        P([("Errors: ", {"size": 11.5, "font": "b", "color": NAVY}), (f"{res['errors']}. ", {"size": 11.5}), ("False positives: ", {"size": 11.5, "font": "b", "color": NAVY}), (f"{res['false_positives']}.", {"size": 11.5})], after=3),
        P([("p50 is the median time, p95 the slowest 5%, in seconds on the live app. The image and voice rows have only 4 and 1 cases.", {"size": 12, "color": MUTED})]),
    ])
    rx = MX + lw_ + 0.2
    rw = CW - lw_ - 0.2
    top = 2.52
    card(s, rx, top, rw, Y1 - top)
    tx(s, rx + 0.18, top + 0.12, rw - 0.36, 0.26, [P([("Read this honestly", {"size": 13, "font": "b", "color": BLUE})])])
    cmp = EVAL / "run-gemini-3.8-flash.md"
    c38 = read_results(cmp) if cmp.exists() else None
    bullets = [P([(f"{n} cases, written by the authors: ", {"size": 11.5, "font": "b", "color": NAVY}), (f"a small test, not a guarantee; real traffic will differ. {ns} scams over 13 types, {ng} genuine.", {"size": 11.5})], bullet="•", after=4)]
    if c38:
        i95 = round(float(c38["latency"]["image"][1].split()[0]))
        bullets.append(P([("Why gemini-3.7-flash: ", {"size": 11.5, "font": "b", "color": NAVY}),
                          (f"an earlier run on gemini-3.8-flash scored {c38['accuracy']} accuracy with {c38['fpr']} false positives and image checks up to {i95} s. Fallback: 2.5-flash.", {"size": 11.5})], bullet="•", after=4))
    else:
        bullets.append(P([("Model: ", {"size": 11.5, "font": "b", "color": NAVY}), ("gemini-3.7-flash, with gemini-2.5-flash as fallback.", {"size": 11.5})], bullet="•", after=4))
    bullets.append(P([("The first run found a real bug ", {"size": 11.5, "font": "b", "color": NAVY}), ("(the model did not know today's date); fixed and re-measured.", {"size": 11.5})], bullet="•"))
    tx(s, rx + 0.18, top + 0.45, rw - 0.36, Y1 - top - 0.5, bullets)
    es = []
    es.append(f"English run ({res['date']}): accuracy {res['accuracy']}, scam recall {res['recall']}, false-positive rate {res['fpr']}, scam type correct {res['type']}, errors {res['errors']} (eval/results/latest.md).")
    if hi:
        es.append(f"Hindi run ({hi['date']}): accuracy {hi['accuracy']}, recall {hi['recall']}, false positives {hi['fpr']}, type correct {hi['type']} (eval/results/run-hindi.md).")
    es.append("Both runs were made against the live app with 50 synthetic cases in eval/cases.jsonl (45 text, 4 image, 1 audio; 26 scams over 13 scam types, 24 genuine), concurrency 1. Latency is wall-clock from request to response. The cases are few and were written by the authors, so real traffic will differ.")
    es.append("The model for these runs was gemini-3.7-flash: eval/results/latest.md is identical to eval/results/run-gemini-3.7-flash-17oct.md, and gemini-3.7-flash is the API default (api/src/index.ts).")
    if c38:
        es.append(f"Earlier comparison on gemini-3.8-flash ({c38['date']}): accuracy {c38['accuracy']}, false-positive rate {c38['fpr']}, image p95 {c38['latency'].get('image', ('', '?', ''))[1]}; that is why the default is gemini-3.7-flash, with gemini-2.5-flash as fallback.")
    es.append("The first benchmark run exposed a real bug (the model did not know today's date and flagged 2026 dates as future); it was fixed and re-measured. This is a small synthetic sanity check, not a guarantee of real-world accuracy.")
    notes(s, " ".join(es))


def slide13(prs):  # Future development
    s = slide_at(prs, 13)
    title_style(s)
    items = [
        ("languages", "More Indian languages", "Tamil, Telugu, Marathi, Bengali and more."),
        ("landmark", "Bank & telecom partners", "Verified reports feed the shared scam list."),
        ("globe", "URL reputation feeds", "Add Google Web Risk next to the isolated scan."),
        ("smartphone", "Android app", "Share SMS and notifications to a check."),
        ("search", "Google Search grounding", "Switch on for link checks once quota allows."),
        ("phone", "Call-recording upload", "Check recorded calls, not just voice notes."),
        ("brain", "Misinformation mode", "The same engine for \"is this forward real?\""),
    ]
    cw, gap, ch = 2.15, 0.2, 1.2
    for i, (ic, t, b) in enumerate(items):
        col, row = i % 4, i // 4
        x, y = MX + col * (cw + gap), Y0 + row * (ch + 0.12)
        card(s, x, y, cw, ch)
        icon_chip(s, ic, x + 0.14, y + 0.14, 0.38)
        tx(s, x + 0.62, y + 0.12, cw - 0.72, 0.42, [P([(t, {"size": 12, "font": "b", "color": NAVY})])], anchor=MSO_ANCHOR.MIDDLE)
        tx(s, x + 0.16, y + 0.62, cw - 0.3, ch - 0.66, [P([(b, {"size": 11.5})])])
    x, y = MX + 3 * (cw + gap), Y0 + (ch + 0.12)
    card(s, x, y, cw, ch, fill=TINT2, line=None)
    tx(s, x + 0.16, y, cw - 0.3, ch, [P([("Plans, not shipped. ", {"size": 11.5, "font": "b", "color": BLUE}), ("Search grounding is built but off by default.", {"size": 11.5, "color": INK})])], anchor=MSO_ANCHOR.MIDDLE)
    yb = Y0 + 2 * ch + 0.12 + 0.14
    card(s, MX, yb, CW, Y1 - yb, fill=NAVY, line=None, radius=0.09)
    tx(s, MX + 0.22, yb + 0.12, CW - 0.44, 0.28, [P([("How it scales", {"size": 13, "font": "b", "color": CYAN})])])
    cols = [
        ("Stateless API", "Cloud Run adds instances (1 to 5 today) as traffic grows; the daily cap is a setting, not a rebuild."),
        ("Managed data", "Firestore, Hosting, Auth and Storage need no servers to run; uploads auto-delete after 7 days."),
        ("Cost stays bounded", "Per-user, per-IP and global limits rise only as budget allows; App Check adds abuse protection next."),
    ]
    w3 = (CW - 0.44) / 3
    for i, (t, b) in enumerate(cols):
        x = MX + 0.22 + i * w3
        tx(s, x, yb + 0.42, w3 - 0.18, Y1 - yb - 0.46, [P([(t, {"size": 12, "font": "b", "color": WHITE})], after=2), P([(b, {"size": 12, "color": rgb("#C9D6EA")})])])
    notes(s, "Roadmap items are plans, not shipped features. Scale story, from what is deployed today (scripts/deploy.sh): API on Cloud Run, stateless, concurrency 20, min 1 and max 5 instances; link scanner scales to zero (max 3); "
             "Firestore for the shared scam list; Firebase Hosting, Auth and Storage; per-user (30 per day), per-IP (60 per minute) and global (3,000 per day) limits; uploads auto-delete after 7 days. "
             "Misinformation mode ('is this forward real?') is an idea for the same engine, not built. Raising the cap or max instances is a configuration change. Hardening ideas: Firebase App Check and a stricter per-device limit (DETAILS.md). Google Search grounding is implemented but off by default because of quota errors on this project.")


def slide14(prs):  # Links
    s = slide_at(prs, 14)
    tb = next(sh for sh in s.shapes if sh.has_text_frame and "Provide links" in sh.text_frame.text)
    tb.left, tb.top, tb.width, tb.height = Inches(0.29), Inches(0.73), Inches(5.75), Inches(3.5)
    values = {"GitHub Public Repository": (REPO_URL, BLUE), "Demo Video Link (3 Minutes)": (VIDEO, AMBER if VIDEO == VIDEO_PLACEHOLDER else BLUE), "Final Product Link": (LIVE_URL, BLUE)}
    for p in tb.text_frame.paragraphs:
        t = p.text.strip()
        if t.startswith("Provide links"):
            for r in p.runs:
                r.font.name, r.font.size = FONT_BOLD, Pt(16)
                r.font.color.rgb = NAVY
        elif t in values:
            val, col = values[t]
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(int(Inches(0.3))))
            pPr.set("indent", str(-int(Inches(0.22))))
            for bu in pPr.findall(qn("a:buSzPts")):
                bu.set("val", "1300")
            for r in p.runs:
                style_run(r, 15, "b", NAVY)
            p.space_after = Pt(22)
            p.add_line_break()
            r = p.add_run()
            r.text = val
            style_run(r, 14, "b" if col == AMBER else None, col)
            if col == AMBER:
                rPr = r._r.get_or_add_rPr()
                rPr.append(rPr.makeelement(qn("a:highlight"), {}))
                rPr.find(qn("a:highlight")).append(rPr.makeelement(qn("a:srgbClr"), {"val": "FFE9A8"}))
                # schema order: fill and highlight must precede the typeface elements
                hl = rPr.find(qn("a:highlight"))
                rPr.remove(hl)
                lat = rPr.find(qn("a:latin"))
                lat.addprevious(hl)
                r2 = p.add_run()
                r2.text = "  PLACEHOLDER: owner to fill in"
                style_run(r2, 11.5, "b", AMBER)
        elif not t:
            for er in p._p.findall(qn("a:endParaRPr")):
                er.set("sz", "800")
            for r in p.runs:
                r.font.size = Pt(8)
    qx, qy, qw, qh = 6.45, 1.0, 3.15, 3.45
    card(s, qx, qy, qw, qh, fill=WHITE, line=BORDER)
    qr = qr_image(LIVE_URL)
    pic(s, qr, qx + 0.5, qy + 0.22, w=2.15)
    tx(s, qx + 0.15, qy + 2.5, qw - 0.3, 0.9, [P([("Scan to try the live app", {"size": 12.5, "font": "b", "color": NAVY})], align=PP_ALIGN.CENTER, after=2),
                                               P([("Paste a suspicious message and see the verdict.", {"size": 12, "color": MUTED})], align=PP_ALIGN.CENTER)])
    yb = 4.62
    card(s, MX, yb, CW, Y1 - yb, fill=TINT2, line=None, radius=0.08)
    pic(s, IMG / "logo-lockup-horizontal.png", MX + 0.2, yb + 0.12, h=0.5)
    tx(s, MX + 3.0, yb, CW - 3.2, Y1 - yb, [P([("Pause. Check. ", {"size": 15, "font": "b", "color": NAVY}), ("Then act.", {"size": 15, "font": "b", "color": BLUE})], align=PP_ALIGN.RIGHT)], anchor=MSO_ANCHOR.MIDDLE)
    if VIDEO == VIDEO_PLACEHOLDER:
        vnote = ("The demo video link is a deliberate placeholder, highlighted in amber; before submitting, re-run the build with DECK_VIDEO_URL=https://... "
                 "(DECK_VIDEO_URL=https://... bash scripts/build-deck.sh) or replace the text '[add the public video link here]' in the PPTX.")
    else:
        vnote = f"Demo video (3 minutes): {VIDEO}"
    notes(s, "GitHub (public): https://github.com/Namans12/scam-shield . Live product: https://project-91260b43-6705-4176-b71.web.app (the QR code points here). " + vnote)


def slide15(prs):  # closing content
    s = slide_at(prs, 15)
    pic(s, IMG / "logo-lockup-horizontal.png", 5 - 1.9, 1.5, w=3.8)
    tx(s, MX, 2.55, CW, 0.6, [P([("Pause. Check. ", {"size": 28, "font": "b", "color": NAVY}), ("Then act.", {"size": 28, "font": "b", "color": BLUE})], align=PP_ALIGN.CENTER)])
    tx(s, 1.4, 3.3, 7.2, 0.5, [P([("Check what you were just sent before you act on it. Free, in English and Hindi.", {"size": 13, "color": MUTED})], align=PP_ALIGN.CENTER)])
    pillars = [("sparkles", "Gemini reads any input"), ("quote", "Evidence for every flag"), ("flag", "A scam becomes action")]
    pw, gap = 2.7, 0.25
    x0 = MX + (CW - (3 * pw + 2 * gap)) / 2
    for i, (ic, t) in enumerate(pillars):
        x = x0 + i * (pw + gap)
        card(s, x, 4.05, pw, 0.8, fill=TINT2, line=None, radius=0.09)
        icon(s, ic, x + 0.2, 4.05 + 0.25, 0.3)
        tx(s, x + 0.65, 4.05, pw - 0.75, 0.8, [P([(t, {"size": 12.5, "font": "b", "color": NAVY})])], anchor=MSO_ANCHOR.MIDDLE)
    tx(s, MX, 5.0, CW, 0.3, [P([(LIVE_URL.replace("https://", ""), {"size": 11, "color": MUTED})], align=PP_ALIGN.CENTER)])
    notes(s, "Closing slide: Scam Shield, 'Pause. Check. Then act.' The live app address is at the bottom. The template's final 'Thank you' slide follows.")


def main(template: str, out: str) -> None:
    res = read_results(EVAL / "latest.md")
    hi_path = EVAL / "run-hindi.md"
    hi = read_results(hi_path) if hi_path.exists() else None
    prs = Presentation(template)
    slide1(prs)
    slide3(prs)
    slide4(prs)
    slide5(prs)
    slide6(prs)
    slide7(prs)
    slide8(prs)
    slide9(prs)
    slide10(prs)
    slide11(prs)
    slide12(prs, res, hi)
    slide13(prs)
    slide14(prs)
    slide15(prs)
    prs.slides[1].notes_slide.notes_text_frame.text = "Organisers' instruction slide, kept exactly as in the template. Nothing to present here; the next slide starts the pitch."
    prs.slides[15].notes_slide.notes_text_frame.text = "Thank you. Team Scam Shield: Naman Shrimal. Live app: https://project-91260b43-6705-4176-b71.web.app ; code: https://github.com/Namans12/scam-shield ."
    cp = prs.core_properties
    cp.title = "Scam Shield - AI Builder Cup pitch deck"
    cp.author = "Scam Shield"
    cp.last_modified_by = "Scam Shield"
    cp.created = cp.modified = datetime.datetime(2026, 10, 18, 0, 0, 0)  # fixed, so rebuilding gives the same file
    prs.save(out)
    print("saved", out)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
