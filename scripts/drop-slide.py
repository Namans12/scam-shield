"""Writes a copy of a PPTX without the given 1-based slide (used to drop the organisers' template note, slide 2, from the submission PDF)."""
import sys
from pptx import Presentation

src, dst, n = sys.argv[1], sys.argv[2], int(sys.argv[3])
prs = Presentation(src)
ids = prs.slides._sldIdLst
sld = list(ids)[n - 1]
prs.part.drop_rel(sld.rId)
ids.remove(sld)
prs.save(dst)
print(f"dropped slide {n}: {len(prs.slides)} slides left")
